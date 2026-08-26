"""M16 前后端契约：采集来源启停、导出、标签和求职阶段。"""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from backend import collection_runs, config, greeting, main, workflow
from backend.db import get_db, init_db, now_iso


class BackendContractTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="boss-copilot-m16-")
        root = Path(self.tmp.name)
        config.DATA_DIR = root
        config.DB_PATH = root / "copilot.db"
        config.COLLECT_RESULT_DIR = root / "job-result"
        init_db()

    def tearDown(self):
        self.tmp.cleanup()

    @staticmethod
    def add_job(key, title="嵌入式工程师", company="示例科技", jd=""):
        ts = now_iso()
        conn = get_db()
        conn.execute(
            "INSERT INTO jobs(job_key,title,company,salary,status,first_seen_at,last_seen_at) "
            "VALUES(?,?,?,'20-30K','active',?,?)", (key, title, company, ts, ts))
        if jd:
            conn.execute(
                "INSERT INTO job_details(job_key,jd,fetched_at) VALUES(?,?,?)",
                (key, jd, ts))
        conn.commit()

    @staticmethod
    def add_run(enabled=True):
        conn = get_db()
        run_id = conn.execute(
            "INSERT INTO collect_runs(kind,params,stats,started_at,finished_at,status,"
            "enabled,phase) VALUES('search','{}','{}',?,?,'succeeded',?,'finished')",
            (now_iso(), now_iso(), 1 if enabled else 0)).lastrowid
        conn.commit()
        return run_id

    def test_disabled_source_hides_only_owned_jobs(self):
        self.add_job("enabled")
        self.add_job("disabled")
        self.add_job("legacy")
        on_run = self.add_run(True)
        off_run = self.add_run(False)
        collection_runs.attach_jobs(on_run, ["enabled"])
        collection_runs.attach_jobs(off_run, ["disabled"])

        visible = {item["job_key"] for item in main.list_jobs(limit=50)["items"]}
        self.assertEqual(visible, {"enabled", "legacy"})
        main.run_enabled(off_run, {"enabled": True})
        visible = {item["job_key"] for item in main.list_jobs(limit=50)["items"]}
        self.assertEqual(visible, {"enabled", "disabled", "legacy"})

    def test_workflow_and_jd_tags_are_exposed_on_list_and_detail(self):
        self.add_job(
            "tagged", jd="团队周末双休，提供五险一金、年终奖和定期体检。")
        state = workflow.set_stage("tagged", "interviewed", True)
        self.assertTrue(state["greeted"])
        self.assertTrue(state["applied"])
        self.assertTrue(state["interviewed"])

        item = main.list_jobs(limit=50)["items"][0]
        self.assertTrue(item["has_weekend"])
        self.assertTrue(item["has_benefits"])
        self.assertTrue(item["interviewed"])
        detail = main.job_detail("tagged")
        self.assertTrue(detail["workflow"]["applied"])

        workflow.set_stage("tagged", "greeted", False)
        self.assertFalse(workflow.get_state("tagged")["interviewed"])

    def test_run_export_contains_owned_jobs(self):
        self.add_job("exported", title="固件工程师", company="导出公司")
        run_id = self.add_run(True)
        collection_runs.attach_jobs(run_id, ["exported"])
        book = load_workbook(collection_runs.export_xlsx(run_id), read_only=True)
        rows = list(book["采集结果"].iter_rows(values_only=True))
        self.assertEqual(rows[0][0:3], ("岗位", "公司", "薪资"))
        self.assertEqual(rows[1][0:2], ("固件工程师", "导出公司"))

    def test_sender_pending_batch_can_be_limited_to_selected_jobs(self):
        self.add_job("selected")
        self.add_job("other")
        resume = get_db().execute(
            "SELECT id,revision FROM resumes WHERE is_default=1").fetchone()
        ts = now_iso()
        for key in ("selected", "other"):
            get_db().execute(
                "INSERT INTO greetings(job_key,variants,chosen,status,created_at,"
                "updated_at,resume_id,resume_revision) VALUES(?,?,'你好','approved',?,?,?,?)",
                (key, '["你好"]', ts, ts, resume["id"], resume["revision"]))
        get_db().commit()
        self.assertEqual(
            [item["job_key"] for item in greeting.pending_batch(["selected"])],
            ["selected"])


if __name__ == "__main__":
    unittest.main()
